"""Tests for the pdf converters in file_converter_util.

Covers the path bug that could write converted files OUTSIDE the user's output directory, and the
new pdf --> docx converter that keeps the text AND the images embedded in the pdf.
"""
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'src'))

# pypdf reads PIL.__version__ at import time, and conftest has replaced PIL with a MagicMock (whose
# __version__ raises AttributeError). Lift the PIL stubs for the duration of these imports, then put
# them back so the other test modules -- which never touch a real image -- keep their stubs.
_PIL_stubs = {name: sys.modules.pop(name)
              for name in ('PIL', 'PIL.Image') if name in sys.modules}

# CI installs pytest and nothing else; skip the whole module rather than abort collection when the
# pdf/docx stack is absent. conftest stubs GUI_util & friends, so the import itself is Tk-free.
try:
    pytest.importorskip('PIL')
    pytest.importorskip('pdfminer')
    pytest.importorskip('docx')
    pytest.importorskip('pypdf')
    pytest.importorskip('striprtf')

    import file_converter_util as fc  # noqa: E402
except BaseException:
    # the real stack is not installed (pytest.importorskip raises Skipped, a BaseException):
    # hand the other test modules their stubs back before bowing out
    sys.modules.update(_PIL_stubs)
    raise
# on success the REAL PIL stays in sys.modules: pypdf decodes the embedded images through it at
# call time, and any other module importing PIL now gets a working one.

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# a real, repo-tracked pdf that embeds 3 images (2 png + 1 jpg)
TIPS_PDF = os.path.join(REPO, 'TIPS', 'TIPS_NLP_Annotator dictionary.pdf')


class TestGetConvertedOutputPath:
    def test_a_sibling_directory_sharing_a_name_prefix_stays_inside_the_output_dir(self):
        """os.path.commonprefix compared STRINGS: 'C:/data/corpus' vs a doc in 'C:/data/corpus2'
        yielded a relative path of '../corpus2/a.pdf', writing the output OUTSIDE outputDir."""
        out = fc.get_converted_output_path(os.path.join('C:', os.sep, 'data', 'corpus2', 'a.pdf'),
                                           os.path.join('C:', os.sep, 'data', 'corpus'),
                                           os.path.join('C:', os.sep, 'out'), '.txt')
        assert os.pardir not in out.split(os.sep)
        assert os.path.basename(out) == 'a.txt'

    def test_a_subdirectory_is_mirrored_under_the_output_dir(self):
        doc = os.path.join('C:', os.sep, 'corpus', 'novels', '1920', 'a.pdf')
        out = fc.get_converted_output_path(doc, os.path.join('C:', os.sep, 'corpus'),
                                           os.path.join('C:', os.sep, 'out'), '.txt')
        assert out == os.path.join('C:', os.sep, 'out', 'novels', '1920', 'a.txt')

    def test_a_file_directly_in_the_input_dir_lands_directly_in_the_output_dir(self):
        out = fc.get_converted_output_path(os.path.join('C:', os.sep, 'corpus', 'a.pdf'),
                                           os.path.join('C:', os.sep, 'corpus'),
                                           os.path.join('C:', os.sep, 'out'), '.txt')
        assert out == os.path.join('C:', os.sep, 'out', 'a.txt')

    def test_no_input_dir_falls_back_to_the_basename(self):
        out = fc.get_converted_output_path(os.path.join('C:', os.sep, 'x', 'y', 'a.pdf'), '',
                                           os.path.join('C:', os.sep, 'out'), '.txt')
        assert out == os.path.join('C:', os.sep, 'out', 'a.txt')

    def test_the_docx_extension_is_used_when_asked_for(self):
        out = fc.get_converted_output_path(os.path.join('C:', os.sep, 'corpus', 'a.pdf'),
                                           os.path.join('C:', os.sep, 'corpus'),
                                           os.path.join('C:', os.sep, 'out'), '.docx')
        assert out.endswith('a.docx')

    def test_an_upper_case_PDF_keeps_its_name(self):
        out = fc.get_converted_output_path(os.path.join('C:', os.sep, 'corpus', 'A.PDF'),
                                           os.path.join('C:', os.sep, 'corpus'),
                                           os.path.join('C:', os.sep, 'out'), '.txt')
        assert os.path.basename(out) == 'A.txt'

    def test_a_dot_in_the_filename_is_not_mistaken_for_the_extension(self):
        out = fc.get_converted_output_path(os.path.join('C:', os.sep, 'c', 'v1.2 draft.pdf'),
                                           os.path.join('C:', os.sep, 'c'),
                                           os.path.join('C:', os.sep, 'out'), '.txt')
        assert os.path.basename(out) == 'v1.2 draft.txt'


class TestGetDocxImageWidth:
    def test_a_pdf_width_is_converted_from_points_to_inches(self):
        from docx.shared import Inches
        assert fc.get_docx_image_width(144) == Inches(2)     # 144 pt = 2 in

    def test_an_over_wide_image_is_capped_to_the_text_column(self):
        from docx.shared import Inches
        assert fc.get_docx_image_width(1440) == Inches(6)    # 20 in -> capped

    def test_a_zero_or_negative_width_falls_back_to_the_cap(self):
        from docx.shared import Inches
        assert fc.get_docx_image_width(0) == Inches(6)
        assert fc.get_docx_image_width(-10) == Inches(6)

    def test_a_non_numeric_width_falls_back_to_the_cap(self):
        from docx.shared import Inches
        assert fc.get_docx_image_width(None) == Inches(6)
        assert fc.get_docx_image_width('wide') == Inches(6)


class _FakeChar:
    """Stands in for a pdfminer LTChar: font name, size and colour are all it is asked for."""

    def __init__(self, fontname='TimesNewRomanPSMT', size=12.0, ncolor=None):
        self.fontname = fontname
        self.size = size
        self.graphicstate = type('gs', (), {'ncolor': ncolor})()


class TestGetCharStyle:
    """A pdf has no bold/italic flag: the typeface lives in the font NAME."""

    def test_bold_is_read_from_the_font_name(self):
        bold, italic, size, colour = fc.get_char_style(_FakeChar('TimesNewRomanPS-BoldMT'))
        assert (bold, italic) == (True, False)

    def test_italic_is_read_from_the_font_name(self):
        bold, italic, size, colour = fc.get_char_style(_FakeChar('TimesNewRomanPS-ItalicMT'))
        assert (bold, italic) == (False, True)

    def test_bold_italic_is_read_from_the_font_name(self):
        bold, italic, size, colour = fc.get_char_style(_FakeChar('TimesNewRomanPS-BoldItalicMT'))
        assert (bold, italic) == (True, True)

    def test_a_plain_font_is_neither(self):
        bold, italic, size, colour = fc.get_char_style(_FakeChar('TimesNewRomanPSMT'))
        assert (bold, italic) == (False, False)

    def test_the_size_comes_through(self):
        assert fc.get_char_style(_FakeChar(size=14.04))[2] == 14.0


class TestGetCharColour:
    def test_rgb_red_is_converted_to_0_255(self):
        assert fc.get_char_colour(_FakeChar(ncolor=[1, 0, 0])) == (255, 0, 0)

    def test_plain_black_is_left_to_the_docx_default(self):
        assert fc.get_char_colour(_FakeChar(ncolor=[0])) is None
        assert fc.get_char_colour(_FakeChar(ncolor=0)) is None

    def test_a_single_value_is_read_as_grayscale(self):
        assert fc.get_char_colour(_FakeChar(ncolor=[0.5])) == (128, 128, 128)

    def test_four_values_are_read_as_cmyk(self):
        assert fc.get_char_colour(_FakeChar(ncolor=[0, 1, 1, 0])) == (255, 0, 0)

    def test_a_missing_or_unusable_colour_is_ignored(self):
        assert fc.get_char_colour(_FakeChar(ncolor=None)) is None
        assert fc.get_char_colour(_FakeChar(ncolor=[1, 2])) is None
        assert fc.get_char_colour(_FakeChar(ncolor=['red', 'green', 'blue'])) is None


class TestGetLineSeparator:
    """A pdf breaks lines wherever the page ran out of room; copied over verbatim those breaks land
    in the middle of sentences."""

    def test_lines_are_joined_with_a_blank(self):
        assert fc.get_line_separator('has then been') == ' '

    def test_a_line_already_ending_in_a_blank_gets_no_second_one(self):
        assert fc.get_line_separator('has then been ') == ''

    def test_an_end_of_line_hyphen_joins_tight(self):
        assert fc.get_line_separator('dictionar-') == ''

    def test_a_dash_standing_on_its_own_is_not_hyphenation(self):
        assert fc.get_line_separator('a word -') == ' '

    def test_an_empty_line_adds_nothing(self):
        assert fc.get_line_separator('') == ''


class TestOCRRequirements:
    """OCR needs FOUR pieces: two Python modules and two binaries the installers cannot bundle.

    The previous OCR code was unreachable (the menu entry ran pdfminer, and image_to_string was
    imported inside another function, so it would have raised NameError on the first call).
    """

    def test_every_missing_piece_is_named_with_the_way_to_install_it(self, monkeypatch):
        monkeypatch.setattr(fc, 'get_tesseract_path', lambda: None)
        monkeypatch.setattr(fc, 'get_poppler_path', lambda: None)
        missing = fc.get_OCR_missing_requirements()
        assert any('esseract' in m for m in missing)
        assert any('poppler' in m for m in missing)
        # each entry has to tell the user what to actually do
        assert all(('pip install' in m or 'brew install' in m or 'http' in m) for m in missing)

    def test_nothing_is_reported_missing_when_all_four_are_present(self, monkeypatch):
        monkeypatch.setattr(fc, 'get_tesseract_path', lambda: '/usr/bin/tesseract')
        monkeypatch.setattr(fc, 'get_poppler_path', lambda: '/usr/bin')
        missing = [m for m in fc.get_OCR_missing_requirements() if 'Python module' not in m]
        assert missing == []

    def test_the_binaries_are_looked_for_beyond_PATH(self, tmp_path, monkeypatch):
        """A Mac app launched from the Finder does not see /opt/homebrew/bin, which is why the old
        code carried a commented-out hard-coded homebrew path."""
        import shutil
        monkeypatch.setattr(shutil, 'which', lambda name: None)
        binary = tmp_path / 'tesseract'
        binary.write_text('')
        assert fc.find_executable('tesseract', [str(binary)]) == str(binary)

    def test_an_absent_binary_is_reported_as_absent(self, tmp_path, monkeypatch):
        import shutil
        monkeypatch.setattr(shutil, 'which', lambda name: None)
        assert fc.find_executable('tesseract', [str(tmp_path / 'nowhere')]) is None


class TestTesseractLanguage:
    def test_the_suite_language_names_map_to_tesseract_codes(self):
        assert fc.tesseract_language_codes['italian'] == 'ita'
        assert fc.tesseract_language_codes['english'] == 'eng'

    def test_an_installed_pack_is_used(self, monkeypatch):
        pytesseract = pytest.importorskip('pytesseract')
        monkeypatch.setattr(fc, 'get_tesseract_path', lambda: None)
        monkeypatch.setattr(pytesseract, 'get_languages', lambda config='': ['eng', 'ita'])
        assert fc.get_tesseract_language('Italian') == ('ita', '')

    def test_a_missing_pack_falls_back_to_english_and_says_so(self, monkeypatch):
        pytesseract = pytest.importorskip('pytesseract')
        monkeypatch.setattr(fc, 'get_tesseract_path', lambda: None)
        monkeypatch.setattr(pytesseract, 'get_languages', lambda config='': ['eng'])
        code, missing = fc.get_tesseract_language('Italian')
        assert code == 'eng'
        assert missing == 'ita'   # so the user can be told WHICH pack to install

    def test_an_unknown_language_does_not_crash(self, monkeypatch):
        monkeypatch.setattr(fc, 'get_tesseract_path', lambda: None)
        code, missing = fc.get_tesseract_language('Klingon')
        assert code in ('eng', 'tlh')


class TestPypdfIsOptional:
    """pypdf must stay OPTIONAL: it is needed only by the pdf --> docx option.

    It was briefly added to IO_libraries_util.install_all_Python_packages, which is a FATAL ERROR
    that exits the Suite and runs at MODULE IMPORT -- so an installation predating pypdf was locked
    out of the csv, docx, rtf and pdf --> txt converters too.
    """

    def test_pypdf_is_not_in_the_fatal_import_gate(self):
        source = open(os.path.join(REPO, 'src', 'file_converter_util.py'), encoding='utf-8').read()
        gate = source.split('install_all_Python_packages(', 1)[1].split(')', 1)[0]
        assert 'pypdf' not in gate

    def test_the_converter_checks_for_pypdf_itself(self):
        assert fc.is_pypdf_available() in (True, False)


@pytest.mark.skipif(not os.path.isfile(TIPS_PDF), reason='TIPS pdf fixture not available')
class TestPdfToDocx:
    def test_the_text_is_extracted(self):
        text = fc.extract_pdf_text(TIPS_PDF)
        assert 'Annotation via Dictionary' in text
        assert len(text) > 1000

    def test_the_layout_keeps_text_and_images_in_reading_order(self):
        pages = fc.get_pdf_layout_items(TIPS_PDF)
        assert len(pages) == 4
        kinds = [kind for page in pages for kind, payload in page]
        assert kinds.count('image') == 3
        assert kinds.count('text') > 20
        # the first page opens with its title, NOT with the picture halfway down it
        assert pages[0][0][0] == 'text'

    def test_the_docx_holds_both_the_text_and_every_image(self, tmp_path):
        out = str(tmp_path / 'sub' / 'rebuilt.docx')
        placed, lost = fc.build_docx_from_pdf(TIPS_PDF, out)
        assert (placed, lost) == (3, 0)
        assert os.path.isfile(out)   # the sub-directory was created for us

        from docx import Document
        document = Document(out)
        assert len(document.inline_shapes) == 3
        text = '\n'.join(p.text for p in document.paragraphs)
        assert 'Annotation via Dictionary' in text
        assert 'IMAGE could not be' not in text

    def test_the_output_is_single_spaced_like_the_pdf(self, tmp_path):
        """A new python-docx document inherits Word's docDefaults -- 10 pt after every paragraph
        and 1.15 line spacing -- which read as double spacing next to a single spaced pdf."""
        out = str(tmp_path / 'spacing.docx')
        fc.build_docx_from_pdf(TIPS_PDF, out)

        from docx import Document
        from docx.shared import Pt
        paragraph_format = Document(out).styles['Normal'].paragraph_format
        assert paragraph_format.space_before == Pt(0)
        assert paragraph_format.space_after == Pt(0)
        assert paragraph_format.line_spacing == 1.0

    def test_the_pdf_formatting_survives_into_the_docx(self, tmp_path):
        out = str(tmp_path / 'styled.docx')
        fc.build_docx_from_pdf(TIPS_PDF, out)

        from docx import Document
        document = Document(out)
        runs = [r for p in document.paragraphs for r in p.runs]
        assert any(r.bold for r in runs)
        assert any(r.italic for r in runs)
        # the red heading of the TIPS file
        assert any(r.font.color is not None and r.font.color.rgb is not None
                   and str(r.font.color.rgb) == 'FF0000' for r in runs)
        # 'Dry September', a book title, is italic in the pdf
        italics = ''.join(r.text for r in runs if r.italic)
        assert 'Dry September' in italics

    def test_the_line_breaks_of_the_pdf_do_not_land_mid_sentence(self, tmp_path):
        out = str(tmp_path / 'joined.docx')
        fc.build_docx_from_pdf(TIPS_PDF, out)

        from docx import Document
        document = Document(out)
        paragraph = [p for p in document.paragraphs if 'third case' in p.text][0]
        assert '\n' not in paragraph.text
        assert 'has then been' in paragraph.text          # was split across two lines
        assert '  ' not in paragraph.text                 # and not joined with a doubled blank

    def test_joining_the_lines_loses_no_text(self, tmp_path):
        import re
        out = str(tmp_path / 'complete.docx')
        fc.build_docx_from_pdf(TIPS_PDF, out)

        from docx import Document

        def norm(s):
            return re.sub(r'\s+', ' ', s).strip()

        pdf_words = norm(fc.extract_pdf_text(TIPS_PDF)).split(' ')
        docx_text = norm('\n'.join(p.text for p in Document(out).paragraphs))
        assert [w for w in pdf_words if w and w not in docx_text] == []

    def test_an_unreadable_pdf_raises_rather_than_writing_a_truncated_file(self, tmp_path):
        broken = tmp_path / 'broken.pdf'
        broken.write_bytes(b'%PDF-1.4 this is not a pdf')
        with pytest.raises(Exception):
            fc.build_docx_from_pdf(str(broken), str(tmp_path / 'broken.docx'))
